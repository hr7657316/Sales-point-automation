from datetime import date

from openpyxl import load_workbook

from sales_points.fit_report import rows_from_grid
from sales_points.quarterly import build_quarterly_report, region_of, score_quarter

GRID = [
    ["Monthly Fit Report", "", ""],
    ["", "", ""],
    ["PATIENT INFO", "PRO", "REP", "TEAM", "INS", "TYPE", "DATE RX REC'D",
     "PATIENT STATUS", "DOS", "PRODUCT", "INSURANCE STATUS", "DATE DME REC'D",
     "SERIAL NUMBER", "URGENCY / INCOMPLETE NOTES"],
    ["JANE DOE 01-01-80", "DR SMITH", "LOPICCOLO (M1-11-69)", "ROYLE (PITT)", "ONECALL/PO",
     "PA WC", "06-01-26", "(8.1) FIT", "A", "TCT-RT KNEE-30 DAY RX DUAL", "BILLED",
     "06-10-26", "TCT-1234", ""],
    ["JANE DOE 01-01-80", "DR SMITH", "LOPICCOLO (M1-11-69)", "ROYLE (PITT)", "ONECALL/PO",
     "PA WC", "06-01-26", "(8.1) FIT", "A", "MZ-RT KNEE(LT) DUAL", "BILLED",
     "06-10-26", "MZ-9999", ""],
    ["JOHN ROE 02-02-70", "DR JONES", "HOUSE WEST (M1-21-2)", "MICH", "AAA",
     "MI AUTO", "06-03-26", "(8.1) FIT", "A", "TCT-LUMBAR-30 DAY RX", "BILLED",
     "06-12-26", "TCT-5555", ""],
]


def test_raw_columns_are_kept():
    rows = rows_from_grid(GRID)
    assert rows[0].raw["SERIAL NUMBER"] == "TCT-1234"
    assert rows[0].raw["DATE DME REC'D"] == "06-10-26"


def test_quarterly_report_has_all_and_team_tabs(tmp_path):
    import csv
    report = tmp_path / "june.csv"
    with open(report, "w", newline="") as handle:
        csv.writer(handle).writerows(GRID)
    out = build_quarterly_report([("JUNE 2026", report)], tmp_path / "q.xlsx", "Q TEST")
    wb = load_workbook(out)
    assert wb.sheetnames == [
        "MASTER-GRAND TOTAL", "MASTER - DME", "MASTER - M1SX", "MASTER - PHARMACY",
        "EAST REGION - GRAND TOTAL", "EAST REGION - DME", "EAST REGION - M1SX",
        "EAST REGION - PHARMACY", "WEST REGION - GRAND TOTAL", "WEST REGION - DME",
        "WEST REGION - M1SX", "WEST REGION - PHARMACY"]
    dme = wb["MASTER - DME"]
    header = [c.value for c in dme[1]]
    assert header[0] == "PATIENT NAME / DOB" and header[14] == "POINT TOTAL"
    points = [dme.cell(row=r, column=15).value for r in (2, 3, 4)]
    assert sorted(points) == [300, 500, 500]          # TCT 300, MZ 500 (Gold Pair excluded), MI Auto TCT 500
    assert dme.cell(row=2, column=14).value in {"TCT-1234", "MZ-9999", "TCT-5555"}
    assert dme.cell(row=6, column=15).value == "TOTAL: 1,300"   # 300 + 500 + 500, no Gold Pair
    notes = " ".join(str(dme.cell(row=r, column=17).value) for r in (2, 3, 4))
    assert "Gold Pair +50 on the rep sheet (not in the RSM total)" in notes
    gt = wb["MASTER-GRAND TOTAL"]
    assert gt["C1"].value == "DME" and gt["E1"].value == 1300
    assert gt["E4"].value == 1300          # no M1Sx or Pharmacy input here
    assert wb["WEST REGION - GRAND TOTAL"]["E1"].value == 500   # HOUSE WEST (M1-21-2)
    assert wb["EAST REGION - GRAND TOTAL"]["E1"].value == 800   # LOPICCOLO (M1-11-69)
    west = wb["WEST REGION - DME"]
    assert west.cell(row=2, column=5).value == "MICH" and west.cell(row=3, column=1).value is None
    east = wb["EAST REGION - DME"]
    assert [east.cell(row=r, column=5).value for r in (2, 3)] == ["ROYLE (PITT)", "ROYLE (PITT)"]


def test_score_quarter_marks_split_and_ffw(tmp_path):
    import csv
    grid = [row[:] for row in GRID]
    grid[3][1] = "DR SMITH *"
    grid[3][2] = "LOPICCOLO (M1-11-69) / HOUSE EAST (M1-21-0)"
    report = tmp_path / "m.csv"
    with open(report, "w", newline="") as handle:
        csv.writer(handle).writerows(grid)
    recs = score_quarter([("JUNE 2026", report)])
    first = [r for r in recs if r["cells"][13] == "TCT-1234"][0]
    assert first["ffw"] == "Y"
    assert "Split:" in first["notes"]
    mz = [r for r in recs if r["cells"][13] == "MZ-9999"][0]
    assert mz["points"] == 500          # no Gold Pair here: the TCT is ancillary
    assert first["sort"][0] == date(2026, 6, 10)


def test_region_comes_from_the_rep_code_list():
    """Allissa 09-23: region is by rep, not by team. Miller (ZARNDT IL) is
    WEST; House East is EAST under both M1-21-0 and the legacy M1-11."""
    from sales_points.quarterly import load_rep_regions
    regions = load_rep_regions()
    assert region_of("MILLER (M1-21-4)", "ZARNDT (IL)", regions) == "WEST"
    assert region_of("CHRISTENSEN (M1-19-20)", "ZARNDT (IL)", regions) == "WEST"
    assert region_of("LOPICCOLO (M1-11-69)", "ROYLE (PITT)", regions) == "EAST"
    assert region_of("ROYLE (M1-11)", "PITT", regions) == "EAST"
    assert region_of("HOUSE WEST (M1-21-2)", "MICH", regions) == "WEST"
    # a split row follows its first rep
    assert region_of("LOPICCOLO (M1-11-69) / HOUSE EAST (M1-21-0)", "", regions) == "EAST"
    # unknown code: fall back to the team
    assert region_of("NEWREP (M1-99-9)", "MICH", regions) == "WEST"


def test_deductions_reduce_the_region_and_master_dme(tmp_path):
    import csv
    report = tmp_path / "june.csv"
    with open(report, "w", newline="") as handle:
        csv.writer(handle).writerows(GRID)
    ded = tmp_path / "ded.csv"
    with open(ded, "w", newline="") as handle:
        csv.writer(handle).writerows([["region", "account", "label", "points"],
                                      ["WEST", "House West", "June honorarium", "500"],
                                      ["EAST", "House East", "July + August honorariums", "1250"]])
    out = build_quarterly_report([("JUNE 2026", report)], tmp_path / "q.xlsx", "Q TEST",
                                 deductions_csv=ded)
    wb = load_workbook(out)
    assert wb["WEST REGION - GRAND TOTAL"]["E1"].value == 500 - 500
    assert wb["EAST REGION - GRAND TOTAL"]["E1"].value == 800 - 1250
    assert wb["MASTER-GRAND TOTAL"]["E1"].value == 1300 - 1750
    west = wb["WEST REGION - DME"]
    cells = [west.cell(row=r, column=15).value for r in range(1, west.max_row + 1)]
    assert "TOTAL: 500" in cells and -500 in cells and "NET TOTAL: 0" in cells
